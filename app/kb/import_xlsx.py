# -*- coding: utf-8 -*-
from __future__ import annotations

from pathlib import Path

import openpyxl

from app.config import XLSX_PATH
from app.constants import LABEL_TO_LANG, LANGS
from app.kb.normalize import cell_str, is_empty_trans
from app.kb.store import Store


def _header_map(row: tuple) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for i, val in enumerate(row):
        t = cell_str(val)
        if not t:
            continue
        if t in {"中文意思", "中文", "中文术语", "中"} and "zh" not in mapping:
            mapping["zh"] = i
            continue
        if t in LABEL_TO_LANG and LABEL_TO_LANG[t] not in mapping:
            mapping[LABEL_TO_LANG[t]] = i
            continue
        if t in {"最近版本", "来源版本", "最早版本"} and "source_version" not in mapping:
            mapping["source_version"] = i
    return mapping


def _rows_from_sheet(ws) -> list[dict[str, str]]:
    it = ws.iter_rows(values_only=True)
    header = next(it, None)
    if not header:
        return []
    mapping = _header_map(header)
    if "zh" not in mapping:
        return []
    out: list[dict[str, str]] = []
    for row in it:
        if not row:
            continue
        zh = cell_str(row[mapping["zh"]] if mapping["zh"] < len(row) else None)
        if not zh:
            continue
        item = {"zh": zh}
        for lang in LANGS:
            idx = mapping.get(lang)
            val = ""
            if idx is not None and idx < len(row):
                val = cell_str(row[idx])
            item[lang] = "" if is_empty_trans(val) else val
        if "source_version" in mapping and mapping["source_version"] < len(row):
            item["source_version"] = cell_str(row[mapping["source_version"]])
        out.append(item)
    return out


def import_xlsx(store: Store, path: Path | None = None, kb_id: str = "pitpat-dict") -> dict[str, int]:
    src = path or XLSX_PATH
    if not src.exists():
        return {"glossary": 0, "tm": 0, "missing_file": 1}

    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    result = {"glossary": 0, "tm": 0, "missing_file": 0}
    try:
        if "术语库" in wb.sheetnames:
            rows = _rows_from_sheet(wb["术语库"])
            result["glossary"] = store.bulk_insert("glossary", rows, kb_id=kb_id)
        if "翻译记忆" in wb.sheetnames:
            rows = _rows_from_sheet(wb["翻译记忆"])
            result["tm"] = store.bulk_insert("tm", rows, kb_id=kb_id)
        # 单 sheet 拆分文件
        if result["glossary"] == 0 and result["tm"] == 0 and wb.sheetnames:
            first = wb[wb.sheetnames[0]]
            rows = _rows_from_sheet(first)
            name = src.name
            if "术语" in name:
                result["glossary"] = store.bulk_insert("glossary", rows, kb_id=kb_id)
            else:
                result["tm"] = store.bulk_insert("tm", rows, kb_id=kb_id)
    finally:
        wb.close()
    return result
