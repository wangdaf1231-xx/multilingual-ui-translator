# -*- coding: utf-8 -*-
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.config import EXPORT_DIR, ensure_dirs
from app.constants import LANG_LABEL, LANGS, SOURCE_LABEL, TABLE_HEADERS
from app.kb.normalize import cell_str
from app.kb.store import Store


def _style_header(ws, ncols: int) -> None:
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    for col in range(1, ncols + 1):
        cell = ws.cell(1, col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"


def export_batch(items: list[dict], checks: list[dict]) -> Path:
    ensure_dirs()
    ts = datetime.now().strftime("%Y%m%d_%H%M")
    path = EXPORT_DIR / f"PitPat翻译_{ts}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "翻译结果"
    ws.append(TABLE_HEADERS)
    for item in items:
        row = [cell_str(item.get("zh")).replace("\n", " ")]
        for k in LANGS:
            row.append(cell_str(item.get(k)).replace("\n", " "))
        ws.append(row)
    _style_header(ws, 8)
    for i, w in enumerate([36, 28, 28, 28, 28, 28, 28, 28], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws2 = wb.create_sheet("翻译库检查")
    ws2.append(["中文", "命中来源", "是否七国齐全", "缺少语言", "补译内容"])
    for c in checks:
        source = SOURCE_LABEL.get(c.get("hit_source", ""), c.get("hit_source", ""))
        missing = c.get("missing") or []
        missing_labels = "、".join(LANG_LABEL[k] for k in missing if k in LANG_LABEL)
        fills = c.get("fills") or {}
        fill_txt = "；".join(f"{LANG_LABEL.get(k, k)}={v}" for k, v in fills.items() if v)
        ws2.append(
            [
                c.get("zh", ""),
                source,
                "是" if c.get("complete") else "否",
                missing_labels,
                fill_txt,
            ]
        )
    _style_header(ws2, 5)
    for i, w in enumerate([36, 16, 14, 22, 40], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    wb.save(path)
    return path


def export_kb(store: Store, kb_id: str) -> Path:
    ensure_dirs()
    ts = datetime.now().strftime("%Y%m%d_%H%M")
    path = EXPORT_DIR / f"知识库_{kb_id}_{ts}.xlsx"
    wb = Workbook()
    for title, table in (("术语库", "glossary"), ("翻译记忆", "tm")):
        data = store.list_entries(table, kb_id=kb_id, limit=100000, offset=0)
        ws = wb.active if title == "术语库" else wb.create_sheet(title)
        if title == "术语库":
            ws.title = title
        headers = ["中文", *[LANG_LABEL[k] for k in LANGS]]
        ws.append(headers)
        for it in data["items"]:
            ws.append([it.get("zh", ""), *[it.get(k, "") for k in LANGS]])
        _style_header(ws, 8)
        for i, w in enumerate([36, 28, 28, 28, 28, 28, 28, 28], 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    wb.save(path)
    return path
