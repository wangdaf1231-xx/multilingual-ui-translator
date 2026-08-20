# -*- coding: utf-8 -*-
"""从 PitPat 全版本翻译 Excel 整理翻译记忆 + 术语库。"""
from __future__ import annotations

import re
from collections import defaultdict
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SRC = Path(r"c:\Users\ruze\Desktop\PitPat-翻译文档.xlsx")
OUT = Path(r"C:\Users\ruze\Desktop\app多国语言翻译助手\翻译库.xlsx")

LANGS = ("en", "fr", "de", "it", "pl", "es", "pt")
LANG_LABEL = {
    "en": "英文",
    "fr": "法语",
    "de": "德语",
    "it": "意大利语",
    "pl": "波兰语",
    "es": "西班牙语",
    "pt": "葡萄牙语",
}

ZH_HEADERS = ("中文意思", "中文", "中")
SKIP_SHEETS = set()  # 历史公益仍纳入，仅作缺口补全
JUNK_ZH = {
    "中文",
    "英语",
    "法语",
    "德语",
    "需求方",
    "英文翻译",
    "意大利语",
    "西班牙语",
    "葡萄牙语",
    "波兰语",
    "备注",
}
JUNK_TRANS_RE = re.compile(r"^\d+\s*行$")
HAS_HAN = re.compile(r"[\u4e00-\u9fff]")
SENTENCE_PUNCT = re.compile(r"[。！？\n]")


def version_key(name: str) -> tuple:
    if name == "历史公益":
        return (0, 0, 0)
    nums = []
    for part in re.split(r"[^\d]+", name):
        if part.isdigit():
            nums.append(int(part))
    while len(nums) < 3:
        nums.append(0)
    return tuple(nums[:3])


def cell_str(v) -> str:
    if v is None:
        return ""
    s = str(v).replace("\xa0", " ").replace("\r\n", "\n").replace("\r", "\n")
    return s.strip()


def is_junk_trans(s: str) -> bool:
    if not s:
        return True
    if s in {"None", "null", "/", "-", "—"}:
        return True
    if JUNK_TRANS_RE.match(s):
        return True
    # 目标语言列里出现整句中文，基本是串列
    if len(HAS_HAN.findall(s)) >= 2:
        return True
    return False


def classify_header_cell(text: str) -> str | None:
    t = cell_str(text)
    if not t:
        return None
    # 先匹配中文列，避免「中文意思」被「意」误判成意大利语
    if t in {"中文意思", "中文", "中"}:
        return "zh"
    if "泰" in t:
        return "skip"
    if t in {"备注"} or t.startswith("完成情况") or t.startswith("最终检验"):
        return "note"
    if t == "特殊要求" or t.startswith("特殊要求"):
        return "req"
    if t == "英文修改":
        return "en_rev"
    if t == "原来":
        return "skip"
    if t in {"英文翻译", "英文", "英语", "英"}:
        return "en"
    if t in {"法语翻译", "法语", "法"}:
        return "fr"
    if t in {"德语翻译", "德语", "德"}:
        return "de"
    if t in {"意大利语翻译", "意大利语", "意"}:
        return "it"
    if t in {"波兰语", "波兰语翻译", "波"}:
        return "pl"
    if t in {"西班牙语翻译", "西班牙语", "西"}:
        return "es"
    if t in {"葡萄牙语翻译", "葡萄牙语", "葡"}:
        return "pt"
    return None


def find_header(rows: list[tuple]) -> tuple[int, dict[str, int]] | None:
    """返回 (header_row_index, role -> col_index)。"""
    best = None
    for i, row in enumerate(rows[:8]):
        mapping: dict[str, int] = {}
        for col, val in enumerate(row):
            role = classify_header_cell(val)
            if role and role not in mapping:
                mapping[role] = col
        if "zh" in mapping and ("en" in mapping or "en_rev" in mapping):
            score = len(mapping)
            if best is None or score > best[0]:
                best = (score, i, mapping)
    if best is None:
        return None
    return best[1], best[2]


def lang_completeness(entry: dict) -> int:
    return sum(1 for k in LANGS if entry.get(k))


def build() -> None:
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    sheet_names = sorted(wb.sheetnames, key=version_key, reverse=True)

    # zh -> entry
    tm: dict[str, dict] = {}
    stats = defaultdict(int)
    per_sheet = []

    for name in sheet_names:
        if name in SKIP_SHEETS:
            continue
        ws = wb[name]
        raw_rows = list(ws.iter_rows(values_only=True))
        if not raw_rows:
            continue
        found = find_header(raw_rows)
        if not found:
            stats["no_header"] += 1
            per_sheet.append((name, 0, 0, "无表头"))
            continue
        header_i, mapping = found
        lang_cols = [k for k in mapping if k in LANGS]
        if lang_cols == ["en"]:
            per_sheet.append((name, 0, 0, "仅英文，已跳过"))
            continue
        added = 0
        scanned = 0
        for row in raw_rows[header_i + 1 :]:
            if not row:
                continue
            zh = cell_str(row[mapping["zh"]] if mapping["zh"] < len(row) else None)
            if not zh or zh in JUNK_ZH or not HAS_HAN.search(zh):
                continue
            scanned += 1
            req = ""
            if "req" in mapping and mapping["req"] < len(row):
                req = cell_str(row[mapping["req"]])
            note = ""
            if "note" in mapping and mapping["note"] < len(row):
                note = cell_str(row[mapping["note"]])

            langs = {}
            for lang in LANGS:
                val = ""
                if lang == "en" and "en_rev" in mapping and mapping["en_rev"] < len(row):
                    val = cell_str(row[mapping["en_rev"]])
                if not val and lang in mapping and mapping[lang] < len(row):
                    val = cell_str(row[mapping[lang]])
                if not is_junk_trans(val):
                    langs[lang] = val
            if not langs:
                continue

            added += 1
            if zh not in tm:
                tm[zh] = {
                    "zh": zh,
                    **{k: langs.get(k, "") for k in LANGS},
                    "latest_version": name,
                    "first_version": name,
                    "versions": [name],
                    "req": req,
                    "note": note,
                    "conflict": "",
                }
            else:
                e = tm[zh]
                if name not in e["versions"]:
                    e["versions"].append(name)
                # 已按新→旧扫描：已有译文不覆盖；只补空
                for k in LANGS:
                    if not e.get(k) and langs.get(k):
                        e[k] = langs[k]
                        e["note"] = (e.get("note") or "") + (
                            f"；{name}补{LANG_LABEL[k]}" if e.get("note") else f"{name}补{LANG_LABEL[k]}"
                        )
                if not e.get("req") and req:
                    e["req"] = req
                e["first_version"] = name  # 更旧
        per_sheet.append((name, scanned, added, ",".join(sorted(k for k in mapping if k in LANGS or k == "en"))))
        stats["rows"] += added

    wb.close()

    # 术语库：2–4 字、至少 6 国译文（相对稳定）；单字仅保留七国齐全
    glossary = []
    for zh, e in tm.items():
        compact = re.sub(r"\s+", "", zh)
        n = lang_completeness(e)
        if SENTENCE_PUNCT.search(zh) or re.search(r"[，、：:（）()【】\[\]{}%*]|\{n\}|\$n", zh):
            continue
        if re.search(r"\d|[xX]{2,}|[Y]{2,}|[Z]{2,}", compact):
            continue
        if not e.get("en") or HAS_HAN.search(e["en"]):
            continue
        if len(compact) == 1 and n == 7:
            glossary.append(e)
            continue
        if 2 <= len(compact) <= 4 and n >= 6:
            glossary.append(e)

    glossary.sort(key=lambda x: (len(x["zh"]), x["zh"]))
    tm_rows = sorted(tm.values(), key=lambda x: (x["latest_version"], x["zh"]))

    write_workbook(tm_rows, glossary, per_sheet)
    print(f"TM {len(tm_rows)}  glossary {len(glossary)}  -> {OUT}")


def style_header(ws, ncols: int) -> None:
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    for col in range(1, ncols + 1):
        cell = ws.cell(1, col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{ws.max_row}"


def set_widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_workbook(tm_rows: list[dict], glossary: list[dict], per_sheet: list) -> None:
    out = openpyxl.Workbook()

    # --- 翻译记忆 ---
    ws = out.active
    ws.title = "翻译记忆"
    headers = [
        "中文",
        "英文",
        "法语",
        "德语",
        "意大利语",
        "波兰语",
        "西班牙语",
        "葡萄牙语",
        "已有语言数",
        "缺失语言",
        "最近版本",
        "最早版本",
        "特殊要求",
        "备注",
    ]
    ws.append(headers)
    missing_fill = PatternFill("solid", fgColor="FFF2CC")
    full_fill = PatternFill("solid", fgColor="E2EFDA")
    for e in tm_rows:
        missing = [LANG_LABEL[k] for k in LANGS if not e.get(k)]
        n = 7 - len(missing)
        row = [
            e["zh"],
            e.get("en", ""),
            e.get("fr", ""),
            e.get("de", ""),
            e.get("it", ""),
            e.get("pl", ""),
            e.get("es", ""),
            e.get("pt", ""),
            n,
            "、".join(missing) if missing else "",
            e["latest_version"],
            e["first_version"],
            e.get("req", ""),
            e.get("note", ""),
        ]
        ws.append(row)
        fill = full_fill if n == 7 else missing_fill
        for col in range(1, 9):
            ws.cell(ws.max_row, col).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(ws.max_row, 9).fill = fill
        if missing:
            ws.cell(ws.max_row, 10).fill = missing_fill
    style_header(ws, len(headers))
    set_widths(ws, [36, 28, 28, 28, 28, 28, 28, 28, 12, 22, 12, 12, 24, 24])

    # --- 术语库 ---
    ws2 = out.create_sheet("术语库")
    g_headers = [
        "中文术语",
        "英文",
        "法语",
        "德语",
        "意大利语",
        "波兰语",
        "西班牙语",
        "葡萄牙语",
        "不可译",
        "已有语言数",
        "缺失语言",
        "来源版本",
        "备注",
    ]
    ws2.append(g_headers)
    for e in glossary:
        missing = [LANG_LABEL[k] for k in LANGS if not e.get(k)]
        n = 7 - len(missing)
        untranslatable = "是" if ("PitPat" in e["zh"] or e.get("en") == e["zh"]) else "否"
        ws2.append(
            [
                e["zh"],
                e.get("en", ""),
                e.get("fr", ""),
                e.get("de", ""),
                e.get("it", ""),
                e.get("pl", ""),
                e.get("es", ""),
                e.get("pt", ""),
                untranslatable,
                n,
                "、".join(missing) if missing else "",
                e["latest_version"],
                "短词自动抽取，请产品核对后作为强制译法",
            ]
        )
        for col in range(1, 9):
            ws2.cell(ws2.max_row, col).alignment = Alignment(wrap_text=True, vertical="top")
    dv = DataValidation(type="list", formula1='"是,否"', allow_blank=False)
    ws2.add_data_validation(dv)
    dv.add("I2:I1048576")
    style_header(ws2, len(g_headers))
    set_widths(ws2, [18, 22, 22, 22, 22, 22, 22, 22, 10, 12, 22, 12, 36])

    # --- 按版本扫描 ---
    ws3 = out.create_sheet("来源扫描")
    ws3.append(["sheet", "识别到的中文行", "写入候选行", "识别到的语言列"])
    for name, scanned, added, langs in per_sheet:
        ws3.append([name, scanned, added, langs])
    style_header(ws3, 4)
    set_widths(ws3, [16, 16, 14, 40])

    # --- 说明 ---
    ws4 = out.create_sheet("整理说明")
    lines = [
        ["项", "说明"],
        ["源文件", str(SRC)],
        ["目标语言", "英 / 法 / 德 / 意 / 波 / 西 / 葡；泰语已丢弃"],
        ["合并规则", "按版本从新到旧扫描。同一中文：新版本译文优先，旧版本只补空缺语言，不覆盖新译法。"],
        ["4.17 起", "开始具备七国列；更早版本多为英/法/德，部分 3.x/4.0/4.1 有意/西/葡"],
        ["缺失语言", "黄色「缺失语言」列。该中文再次使用时再补全，不要用旧版不完整译文冒充七国完整。"],
        ["英文修改", "4.0 / 4.1 若有「英文修改」列，优先于「英文翻译」"],
        ["术语库", "2–4 字且至少已有 6 国译文的短词候选，必须产品核对。新版本优先可能导致个别词被较差译法覆盖（如关注=followed）"],
        ["未收录", "无中文、无汉字、表头行、译文为「39行」一类占位、七国全空的行"],
        ["历史公益", "作为最旧来源，仅在其它版本缺译文时补英/法/德"],
    ]
    for line in lines:
        ws4.append(line)
    style_header(ws4, 2)
    set_widths(ws4, [16, 90])
    for r in range(2, ws4.max_row + 1):
        ws4.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
        ws4.row_dimensions[r].height = 32

    OUT.parent.mkdir(parents=True, exist_ok=True)
    out.save(OUT)


if __name__ == "__main__":
    build()
